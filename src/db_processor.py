import logging
import openpyxl
from openpyxl.styles import PatternFill
import pandas as pd
from typing import Dict, List, Tuple, Optional

logger = logging.getLogger(__name__)

# --- CONFIGURATION: COLUMN NAMES ---
KEY_COL_SIDE = "Side"
KEY_COL_DD = "DD_No"
KEY_COL_OUT = "OUT"

DATA_COL_KAMPAN = "Kampan"
DATA_COL_PS = "PS"
DATA_COL_SIDE = "Side"
DATA_COL_DD = "DD"
# -----------------------------------


def build_coordinate_map(key_df: pd.DataFrame) -> Dict[int, List[Tuple[str, int]]]:
    """Builds a mapping dictionary strictly from the configured output column.

    Iterates through the provided key DataFrame and maps values from the
    configured output column to the corresponding side and wire number
    based on the global configuration constants.

    Args:
        key_df: Pandas DataFrame containing mapping key data.

    Returns:
        A dictionary mapping physical coordinates to a list of possible
        (Side, wire_number) tuples.
    """
    coord_map = {}

    for idx, row in key_df.iterrows():
        try:
            side = str(row[KEY_COL_SIDE]).strip()
            dd_no = int(row[KEY_COL_DD])

            if pd.isna(row[KEY_COL_OUT]):
                continue

            out_val = int(row[KEY_COL_OUT])

            if out_val not in coord_map:
                coord_map[out_val] = []

            if (side, dd_no) not in coord_map[out_val]:
                coord_map[out_val].append((side, dd_no))
        except (ValueError, KeyError, TypeError) as e:
            logger.debug(
                "Skipping row %s in key_df due to invalid data or missing columns: %s",
                idx,
                e,
            )
            continue

    return coord_map


def _find_context_sides(
    rows: List[int], current_idx: int, existing_sides: Dict[int, str]
) -> Tuple[Optional[str], Optional[str]]:
    """Finds the nearest defined sides looking forward and backward.

    Args:
        rows: List of row indices in the current block.
        current_idx: The index of the current row being processed.
        existing_sides: Dictionary mapping row indices to known sides.

    Returns:
        A tuple containing (lookbehind_side, lookahead_side).
    """
    lookahead_side = None
    for look_r in rows[current_idx + 1 :]:
        if look_r in existing_sides:
            lookahead_side = existing_sides[look_r]
            break

    lookbehind_side = None
    for look_r in reversed(rows[:current_idx]):
        if look_r in existing_sides:
            lookbehind_side = existing_sides[look_r]
            break

    return lookbehind_side, lookahead_side


def _resolve_ambiguous_side(
    options: List[Tuple[str, int]],
    lookbehind: Optional[str],
    lookahead: Optional[str],
    streak_count: int,
    ring_next: Dict[str, str],
) -> Tuple[Tuple[str, int], bool]:
    """Selects the best fitting side from available options based on context.

    Args:
        options: List of possible (Side, DD) combinations.
        lookbehind: The side identified in previous rows.
        lookahead: The side identified in upcoming rows.
        streak_count: Number of consecutive identical coordinates.
        ring_next: Mapping dictionary for natural side progression.

    Returns:
        A tuple containing the chosen (Side, DD) combination and a boolean
        flag indicating if the decision was uncertain.
    """
    if len(options) == 1:
        return options[0], False

    chosen = None

    if streak_count == 0:
        if lookbehind:
            for opt in options:
                if opt[0] == lookbehind:
                    chosen = opt
                    break
        if chosen is None and lookahead:
            for opt in options:
                if opt[0] == lookahead:
                    chosen = opt
                    break
    else:
        if lookahead:
            for opt in options:
                if opt[0] == lookahead:
                    chosen = opt
                    break
        if chosen is None and lookbehind:
            next_side = ring_next.get(lookbehind)
            for opt in options:
                if opt[0] == next_side:
                    chosen = opt
                    break
        if chosen is None and lookbehind:
            for opt in options:
                if opt[0] == lookbehind:
                    chosen = opt
                    break

    if chosen is None:
        return options[0], True

    return chosen, False


def process_rust_data_safely(
    input_path: str, output_path: str, coord_map: Dict[int, List[Tuple[str, int]]]
) -> None:
    """Fixes coordinates directly in the Excel file using openpyxl.

    Updates target side and wire number cells in-place only for rows where
    the side value is missing. Uses dynamic anchor tracking (lookbehind and
    lookahead) combined with a local streak counter for consecutive ambiguous
    points. Uses global constants for column names to allow easy schema updates.

    Args:
        input_path: Path to the input Excel file.
        output_path: Path to save the processed Excel file.
        coord_map: Dictionary generated by the build_coordinate_map function.
    """
    logger.info("Loading workbook with openpyxl to preserve formatting.")
    wb = openpyxl.load_workbook(input_path)

    sheet_name = "RustDD_Data"
    if sheet_name not in wb.sheetnames:
        logger.error("Data sheet not found in the workbook.")
        return

    sheet = wb[sheet_name]
    warning_fill = PatternFill(
        start_color="FFFF00", end_color="FFFF00", fill_type="solid"
    )

    header_row = 1
    col_indices = {}
    required_data_cols = [DATA_COL_KAMPAN, DATA_COL_PS, DATA_COL_SIDE, DATA_COL_DD]

    for col in range(1, sheet.max_column + 1):
        val = sheet.cell(row=header_row, column=col).value
        if val in required_data_cols:
            col_indices[val] = col

    if not all(k in col_indices for k in required_data_cols):
        logger.error("Could not find all required columns: %s", required_data_cols)
        return

    kampan_col = col_indices[DATA_COL_KAMPAN]
    ps_col = col_indices[DATA_COL_PS]
    side_col = col_indices[DATA_COL_SIDE]
    dd_col = col_indices[DATA_COL_DD]
    ring_next = {"s1": "s2", "s2": "s3", "s3": "s4", "s4": "s5", "s5": "s6", "s6": "s1"}

    blocks = {}
    for row in range(header_row + 1, sheet.max_row + 1):
        k_val = sheet.cell(row=row, column=kampan_col).value
        p_val = sheet.cell(row=row, column=ps_col).value
        block_key = (k_val, p_val)
        if block_key not in blocks:
            blocks[block_key] = []
        blocks[block_key].append(row)

    for block_key, rows in blocks.items():
        row_vals = {}
        existing_sides = {}

        for r in rows:
            try:
                row_vals[r] = int(sheet.cell(row=r, column=dd_col).value)
            except (ValueError, TypeError):
                pass

            current_side_val = sheet.cell(row=r, column=side_col).value
            if current_side_val and str(current_side_val).strip():
                existing_sides[r] = str(current_side_val).strip()
            else:
                val = row_vals.get(r)
                if val in coord_map and len(coord_map[val]) == 1:
                    existing_sides[r] = coord_map[val][0][0]

        last_val = None
        streak_count = 0

        for i, r in enumerate(rows):
            current_side_val = sheet.cell(row=r, column=side_col).value
            if current_side_val and str(current_side_val).strip():
                continue

            val = row_vals.get(r)
            if val not in coord_map:
                continue

            if val == last_val:
                streak_count += 1
            else:
                streak_count = 0
                last_val = val

            side_cell = sheet.cell(row=r, column=side_col)
            dd_cell = sheet.cell(row=r, column=dd_col)
            options = coord_map[val]

            lookbehind, lookahead = _find_context_sides(rows, i, existing_sides)

            chosen, uncertain = _resolve_ambiguous_side(
                options, lookbehind, lookahead, streak_count, ring_next
            )

            existing_sides[r] = chosen[0]
            side_cell.value = chosen[0]
            dd_cell.value = chosen[1]

            if uncertain:
                side_cell.fill = warning_fill
                dd_cell.fill = warning_fill
                logger.warning(
                    "Lost context at row %s (Block=%s). Forced assignment to %s.",
                    r,
                    block_key,
                    chosen,
                )

    logger.info("Saving workbook to %s", output_path)
    wb.save(output_path)
