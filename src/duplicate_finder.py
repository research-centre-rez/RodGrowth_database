import logging
import openpyxl
from openpyxl.styles import PatternFill
import re
from typing import Dict, List, Tuple, Optional

logger = logging.getLogger(__name__)

# --- CONFIGURATION ---
DATA_COL_PS = "PS"
# ---------------------


def extract_base_ps(ps_name: str) -> str:
    """Extracts the base pattern sample name.

    Uses a regular expression to capture the core identifier, typically
    consisting of letters followed by numbers (e.g., 'XX01'). Ignores any
    subsequent characters like underscores or descriptive text.

    Args:
        ps_name: The raw pattern sample string.

    Returns:
        The extracted base string, or the original string if no match is found.
    """
    if not ps_name:
        return ""

    ps_str = str(ps_name).strip()
    match = re.match(r"^([a-zA-Z]+\d+)", ps_str)

    if match:
        return match.group(1)

    return ps_str


def _find_column_index(
    sheet: openpyxl.worksheet.worksheet.Worksheet, header_row: int, column_name: str
) -> Optional[int]:
    """Finds the column index for a given header name.

    Args:
        sheet: The openpyxl worksheet object.
        header_row: The index of the header row.
        column_name: The name of the column to find.

    Returns:
        The integer column index, or None if not found.
    """
    for col in range(1, sheet.max_column + 1):
        if sheet.cell(row=header_row, column=col).value == column_name:
            return col
    return None


def highlight_ps_variants(input_path: str, output_path: str) -> None:
    """Identifies and highlights variant PS names in an Excel workbook.

    Reads the specified column, groups cells by their base PS name, and applies
    a highlight to cells that contain additional text beyond the base name,
    signifying they are duplicates or variants.

    Args:
        input_path: Path to the input Excel file.
        output_path: Path to save the processed Excel file.
    """
    logger.info("Loading workbook to highlight PS variants.")
    wb = openpyxl.load_workbook(input_path)

    sheet_name = "RustDD_Data"
    if sheet_name not in wb.sheetnames:
        logger.error("Sheet %s not found in the workbook.", sheet_name)
        return

    sheet = wb[sheet_name]
    variant_fill = PatternFill(
        start_color="FFA500", end_color="FFA500", fill_type="solid"
    )

    header_row = 1
    ps_col_idx = _find_column_index(sheet, header_row, DATA_COL_PS)

    if not ps_col_idx:
        logger.error("Could not find column: %s", DATA_COL_PS)
        return

    base_groups: Dict[str, List[Tuple[int, str]]] = {}

    for row in range(header_row + 1, sheet.max_row + 1):
        cell_val = sheet.cell(row=row, column=ps_col_idx).value
        if not cell_val:
            continue

        original_str = str(cell_val).strip()
        base_name = extract_base_ps(original_str)

        if base_name not in base_groups:
            base_groups[base_name] = []
        base_groups[base_name].append((row, original_str))

    variants_found = 0

    for base_name, members in base_groups.items():
        if len(members) > 1:
            for row_idx, original_str in members:
                if len(original_str) > len(base_name):
                    sheet.cell(row=row_idx, column=ps_col_idx).fill = variant_fill
                    variants_found += 1
                    logger.debug(
                        "Highlighted variant at row %s: %s", row_idx, original_str
                    )

    logger.info(
        "Found and highlighted %s variants. Saving to %s", variants_found, output_path
    )
    wb.save(output_path)
