import pandas as pd
import pytest
from pathlib import Path
import openpyxl
from unittest.mock import patch

# Import the functions from your src modules
# Uprav cesty importu podle skutečné struktury tvého projektu
from src.db_processor import build_coordinate_map, process_rust_data_safely
from src.duplicate_finder import extract_base_ps, highlight_ps_variants
from src.main import main


def test_extract_base_ps():
    """Tests the regular expression logic for extracting base PS names."""
    assert extract_base_ps("XX01_naklon") == "XX01"
    assert extract_base_ps("XX01") == "XX01"
    assert extract_base_ps("S14_oprava_v2") == "S14"
    assert extract_base_ps("Kampan02") == "Kampan02"
    assert extract_base_ps("") == ""
    assert extract_base_ps(None) == ""


def test_build_coordinate_map():
    """Tests the creation of the coordinate mapping dictionary."""
    data = {
        "Side": ["s1", "s2", "s3", "s1"],
        "DD_No": [1, 1, 2, 2],
        "OUT": [148, 148, 149, None],  # One shared corner, one unique, one NaN
    }
    df = pd.DataFrame(data)

    coord_map = build_coordinate_map(df)

    assert 148 in coord_map
    assert ("s1", 1) in coord_map[148]
    assert ("s2", 1) in coord_map[148]
    assert len(coord_map[148]) == 2

    assert 149 in coord_map
    assert ("s3", 2) in coord_map[149]
    assert len(coord_map[149]) == 1


@pytest.fixture
def sample_excel_file(tmp_path: Path) -> Path:
    """Fixture to create a temporary Excel file for integration testing."""
    file_path = tmp_path / "test_data.xlsx"
    wb = openpyxl.Workbook()

    # Create main data sheet
    ws_data = wb.active
    ws_data.title = "RustDD_Data"
    headers = ["Kampan", "PS", "Side", "DD"]
    ws_data.append(headers)

    # Add some mock data (XX01 is normal, XX01_variant is a duplicate)
    ws_data.append([1, "XX01", "", 148])
    ws_data.append([1, "XX01", "s1", 149])
    ws_data.append([1, "XX01_variant", "s1", 148])

    wb.save(file_path)
    return file_path


def test_highlight_ps_variants(sample_excel_file: Path):
    """Tests if variants are correctly identified and highlighted in Excel."""
    highlight_ps_variants(str(sample_excel_file), str(sample_excel_file))

    wb = openpyxl.load_workbook(sample_excel_file)
    ws = wb["RustDD_Data"]

    # Row 2 (XX01) should NOT be highlighted
    fill_normal = ws.cell(row=2, column=2).fill
    assert fill_normal.start_color.index == "00000000"  # Default transparent

    # Row 4 (XX01_variant) SHOULD be highlighted orange (FFA500)
    fill_variant = ws.cell(row=4, column=2).fill
    assert fill_variant.start_color.index == "00FFA500"


@patch("src.main.process_rust_data_safely")
@patch("src.main.highlight_ps_variants")
@patch("src.main.build_coordinate_map")
@patch("pandas.read_excel")
def test_main_cli(
    mock_read_excel, mock_build_map, mock_highlight, mock_process, tmp_path: Path
):
    """Tests the main CLI pipeline execution flow and argument parsing."""
    input_file = tmp_path / "input.xlsx"
    output_file = tmp_path / "output.xlsx"
    key_file = tmp_path / "key.xlsx"

    # Create fake files so the path.exists() checks pass
    input_file.touch()
    key_file.touch()

    test_args = [
        "main.py",
        "--input",
        str(input_file),
        "--output",
        str(output_file),
        "--key",
        str(key_file),
    ]

    with patch("sys.argv", test_args):
        main()

    # Verify that all steps in the pipeline were called correctly
    mock_read_excel.assert_called_once_with(key_file, sheet_name="DD_key")
    mock_build_map.assert_called_once()
    mock_process.assert_called_once_with(
        str(input_file), str(output_file), mock_build_map.return_value
    )
    mock_highlight.assert_called_once_with(str(output_file), str(output_file))
